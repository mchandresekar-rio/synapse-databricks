from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = (
    ROOT
    / "skills"
    / "synapse-to-azure-databricks"
    / "docs"
    / "synapse-to-azure-databricks-prerequisites.xlsx"
)

NAVY = "1F4E78"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")

STATUSES = [
    "Not started",
    "In progress",
    "Pending approval",
    "Blocked",
    "Not applicable",
    "Verified",
    "Completed",
    "Failed",
]
SECURITY_DISPOSITIONS = ["Map", "Redesign", "Retain temporarily", "Remove", "Blocker"]


def add_table(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[object]],
    widths: list[int],
    table_name: str,
):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D9D9D9"))
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 38
    table = Table(
        displayName=table_name,
        ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    return sheet


def add_list(sheet, column: int, source_column: str, count: int, rows: int):
    validation = DataValidation(
        type="list",
        formula1=f"'Lists'!${source_column}$2:${source_column}${count + 1}",
        allow_blank=True,
    )
    validation.error = "Select a value from the list."
    validation.errorTitle = "Invalid value"
    sheet.add_data_validation(validation)
    validation.add(
        f"{sheet.cell(2, column).coordinate}:{sheet.cell(rows + 1, column).coordinate}"
    )


def add_status_colors(sheet, column: int, rows: int):
    letter = get_column_letter(column)
    colors = {
        "Verified": GREEN,
        "Completed": GREEN,
        "Blocked": RED,
        "Failed": RED,
        "Pending approval": AMBER,
        "In progress": BLUE,
        "Not applicable": GRAY,
    }
    for status, color in colors.items():
        sheet.conditional_formatting.add(
            f"{letter}2:{letter}{rows + 1}",
            FormulaRule(
                formula=[f'{letter}2="{status}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )


def build_workbook() -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.properties.title = "Synapse to Azure Databricks customer setup"
    workbook.properties.subject = "Customer-admin setup before migration user runs the skill"
    workbook.properties.creator = "Skills for Databricks"

    start = workbook.create_sheet("Start Here")
    start.sheet_view.showGridLines = False
    start.column_dimensions["A"].width = 28
    start.column_dimensions["B"].width = 115
    start.merge_cells("A1:B1")
    start["A1"] = "Synapse to Azure Databricks - Customer Setup and Handoff"
    start["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    start["A1"].fill = PatternFill("solid", fgColor=NAVY)
    start["A1"].alignment = Alignment(horizontal="center", vertical="center")
    start.row_dimensions[1].height = 30
    guidance = [
        ("Customer administrator", "Completes only the privileged setup in Customer Setup: identity onboarding, source access or exports, target bootstrap, storage, compute, and scoped grants."),
        ("Migration user", "After customer setup is accepted, the separate named migration user runs the GitHub Copilot skill from the customer VDI and performs discovery, inventory, conversion, deployment, and validation."),
        ("No admin access", "The migration user is not an Azure, Synapse, Databricks account, workspace, or metastore administrator. Catalogs, schemas, storage credentials, external locations, and compute are pre-created by customer administrators."),
        ("Lakebridge gate", "Lakebridge is verified before assessment. If it is missing or unusable, the skill stops and asks whether to install it, use the desktop installer, explicitly approve governed fallback, or stop."),
        ("How to use", "Customer administrators complete Customer Setup in order, replace each placeholder, record the actual action used, attach evidence, and obtain handoff acceptance before the migration user starts."),
        ("Security migration", "After technical schema validation, the migration user uses the skill to build Security Mapping. Customer identity and platform administrators then create or reuse principals, apply approved grants, and revoke temporary migration access."),
        ("Safety", "Never place passwords, tokens, client secrets, SAS values, connection strings, or private keys in this workbook."),
    ]
    for row_index, (label, value) in enumerate(guidance, start=3):
        start.cell(row_index, 1, label)
        start.cell(row_index, 2, value)
        start.cell(row_index, 1).fill = PatternFill("solid", fgColor=BLUE)
        for cell in start[row_index]:
            cell.font = Font(name="Arial", size=10, bold=cell.column == 1, color=NAVY if cell.column == 1 else "000000")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        start.row_dimensions[row_index].height = 44

    setup_rows = [
        ["SET-01", 10, "Identity administrator", "Create or onboard the named migration-user account.", "Create or invite <migration-user-upn> in the customer tenant using the approved employee, guest, or vendor identity process. Enforce MFA and conditional access.", "<migration-user-upn>; <identity-type>; <expiry-or-review-date>", "", "Named account exists, MFA works, and identity object ID is recorded.", "", "Pending approval", "", "", ""],
        ["SET-02", 20, "Customer desktop administrator", "Prepare the customer VDI for the migration user.", "Install approved VS Code, Git, Python, Java, Azure CLI, and Databricks CLI. Enable GitHub Copilot. Validate: code --version; git --version; python --version; java -version; az version; databricks version", "<VDI-name>; <approved-software-versions>; <migration-user-upn>", "", "Version output, VDI owner, persistence location, and approved outbound endpoints.", "", "Not started", "", "", ""],
        ["SET-LB", 25, "Customer desktop and Databricks administrators", "Install and verify Databricks Labs Lakebridge before migration assessment.", "Verify Python 3.10.1-3.14.x, Java 21+ when Morpheus is in scope, Databricks CLI authentication, and approved access to GitHub, Maven Central, and PyPI. After software approval run: databricks labs install lakebridge [--profile <profile_name>]. Verify: databricks labs lakebridge --help. Treat databricks labs lakebridge install-transpile and --include-llm-transpiler true as separate approval-gated installations.", "<Databricks-profile-name>; <Lakebridge-version>; <Python-version>; <Java-version>; <software-approval-id>", "", "Lakebridge help/version evidence, authenticated profile, prerequisite versions, approved network endpoints, and separate transpiler/LLM decision.", "", "Not started", "", "", "If Lakebridge is missing, the skill must ask before installation and must not silently choose governed fallback."],
        ["SET-03", 30, "GitHub and repository administrator", "Grant GitHub Copilot and repository access.", "Assign the required GitHub Copilot seat to <migration-user-upn>. Add the user to <repository-team> with the minimum repository permission needed to create migration artifacts and changes.", "<migration-user-upn>; <GitHub-organization>; <repository>; <repository-team>; <permission>", "", "Copilot sign-in and repository clone or write test succeeds.", "", "Pending approval", "", "", ""],
        ["SET-04", 40, "Synapse owner", "Provide complete source exports or approved read-only discovery access.", "Preferred export: provide all in-scope SQL definitions/projects, dependencies, Spark artifacts, pipelines, triggers, datasets, linked-service metadata with secrets redacted, runtime settings, and file hashes. If live discovery is required, assign only the approved Synapse read role on <workspace>.", "<workspace>; <databases>; <schemas>; <artifact-types>; <export-folder>; <migration-user-upn>", "", "Export counts and SHA-256 hashes reconcile, or read-only list/get probes succeed.", "", "Pending approval", "", "", "No source publish, execute, update, or delete permission."],
        ["SET-05", 50, "Azure subscription or resource owner", "Grant Azure control-plane read access only when live discovery requires it.", "Azure portal > <workspace-resource> > Access control (IAM) > Add role assignment > Reader > assign <migration-user-upn>. Scope the assignment to the workspace resource or narrowest approved parent.", "<workspace-resource-id>; <migration-user-upn>; <expiry>", "", "Role assignment ID and successful read-only ARM query.", "", "Not applicable", "", "", ""],
        ["SET-06", 60, "Synapse SQL database administrator", "Grant SQL metadata and bounded validation access.", "Run in each approved database: CREATE USER [<migration-user-upn>] FROM EXTERNAL PROVIDER; GRANT CONNECT TO [<migration-user-upn>]; GRANT VIEW DEFINITION TO [<migration-user-upn>]; Add GRANT SELECT ON SCHEMA::[<schema>] only for schemas required by the approved validation plan.", "<migration-user-upn>; <database-list>; <validation-schema-list>; <expiry>", "", "Executed statements, database users, effective permissions, and successful metadata query.", "", "Pending approval", "", "", "Do not grant SQL control, alter, create, insert, update, delete, or execute unless separately approved."],
        ["SET-07", 70, "Azure storage owner", "Grant source storage read access only when required by migration or validation.", "Assign Storage Blob Data Reader to <migration-user-object-id> on the exact approved container or path. Apply ADLS ACL read and execute permissions on only the required directories.", "<migration-user-object-id>; <storage-account>; <container>; <path-list>; <expiry>", "", "Azure role assignment ID, ACL evidence, and successful bounded read probe.", "", "Not applicable", "", "", "No write or delete permission on source storage."],
        ["SET-08", 80, "Databricks account or workspace administrator", "Add the migration user to the non-production Databricks workspace without admin entitlement.", "Account or workspace console > User management > assign <migration-user-upn> to <workspace-id>. Enable Workspace access only. Do not enable workspace admin.", "<migration-user-upn>; <workspace-id>; <workspace-url>; <expiry>", "", "Workspace assignment and non-admin entitlement evidence.", "", "Pending approval", "", "", ""],
        ["SET-09", 90, "Metastore administrator or catalog owner", "Create target catalogs, schemas, owners, and managed-storage boundaries.", "After collision checks run approved DDL: SHOW CATALOGS LIKE '<catalog>'; CREATE CATALOG `<catalog>`; CREATE SCHEMA `<catalog>`.`<schema>`; repeat for the approved schema list. Apply the approved initial owner. Add MANAGED LOCATION only when required by the target design.", "<catalog>; <schema-list>; <managed-location-if-approved>; <initial-owner>", "", "Statement IDs, object IDs, definitions, locations, and owner evidence.", "", "Pending approval", "", "", "Migration user does not receive CREATE CATALOG, CREATE SCHEMA, or metastore administration."],
        ["SET-10", 100, "Azure storage and Databricks metastore administrators", "Create target storage credentials and external locations.", "Using approved IaC or Catalog Explorer, create <storage-credential> for <workload-identity> and <external-location> for the exact <abfss-path>. Configure Azure RBAC and ADLS ACLs for that identity and test access.", "<workload-identity>; <storage-account>; <container>; <path>; <storage-credential>; <external-location>", "", "Object IDs, credential owner, path boundary, Azure RBAC/ACL evidence, and access test.", "", "Pending approval", "", "", "Migration user does not own or MANAGE the storage credential or external location."],
        ["SET-11", 110, "Databricks platform administrator", "Create or select the SQL warehouse and compute policy.", "Databricks SQL > SQL Warehouses: create or select <warehouse>; configure approved size, network, channel, tags, auto-stop, concurrency, and cost limits. Grant CAN USE to <migration-user-principal>. If Spark is required, grant CAN USE only on <compute-policy>.", "<warehouse>; <warehouse-id>; <compute-policy-if-required>; <migration-user-principal>; <auto-stop>; <cost-limit>", "", "Warehouse and policy IDs, ACL evidence, configuration export, and stopped-state proof.", "", "Pending approval", "", "", ""],
        ["SET-12", 120, "Catalog or schema owner", "Grant exact non-admin Unity Catalog privileges required for schema migration.", "Run for each approved schema: GRANT USE CATALOG ON CATALOG `<catalog>` TO `<migration-user-principal>`; GRANT USE SCHEMA, CREATE TABLE ON SCHEMA `<catalog>`.`<schema>` TO `<migration-user-principal>`; add CREATE FUNCTION or CREATE MATERIALIZED VIEW only for object types in scope; grant SELECT only on existing objects referenced by migrated definitions.", "<catalog>; <schema-list>; <migration-user-principal>; <required-object-types>; <referenced-existing-objects>; <expiry>", "", "Executed grants and effective privilege report for the migration user.", "", "Pending approval", "", "", "Do not grant ALL PRIVILEGES, MANAGE, CREATE CATALOG, CREATE SCHEMA, broad MODIFY, or access outside approved schemas."],
        ["SET-13", 130, "External-location owner", "Allow external-table registration only on approved target paths.", "If the migration includes external tables, run: GRANT CREATE EXTERNAL TABLE ON EXTERNAL LOCATION `<external-location>` TO `<migration-user-principal>`; confirm CREATE TABLE is already limited to the approved schema.", "<external-location>; <migration-user-principal>; <approved-paths>; <expiry>", "", "Effective external-location privilege and path-boundary evidence.", "", "Not applicable", "", "", "Do not grant external-location ownership, MANAGE, WRITE FILES, or broader locations for schema-only migration."],
        ["SET-14", 140, "Security and Databricks administrators", "Enable OAuth sign-in and verify the migration user's effective access.", "On the customer VDI, migration user runs: az login; databricks auth login --host <workspace-url>; databricks auth profiles. Customer admin verifies no personal access token is required and reviews effective source, workspace, warehouse, schema, and external-location permissions.", "<workspace-url>; <Databricks-profile-name>; <migration-user-upn>", "", "Successful OAuth login, identity confirmation, positive access probes, and negative admin/out-of-scope probes.", "", "Pending approval", "", "", ""],
        ["SET-15", 150, "Migration lead, security owner, and platform owner", "Approve handoff to the migration user.", "Review every prior setup row, including SET-LB. Confirm all required rows are Completed or Verified, every temporary permission has an expiry, no admin role is assigned, and rollback/support contacts are available. Record the handoff approval ID.", "<handoff-approval-id>; <migration-user-upn>; <support-contacts>; <expiry-register>", "", "Signed handoff authorizing the migration user to start the GitHub Copilot skill.", "", "Pending approval", "", "", "The migration user starts discovery and migration only after this row is approved."],
    ]
    platforms = [
        "Microsoft Entra admin center",
        "Customer VDI and software management",
        "Customer VDI, Databricks CLI, and Lakebridge",
        "GitHub organization and repository",
        "Synapse Studio or customer export process",
        "Azure portal - Access control (IAM)",
        "Synapse serverless or dedicated SQL endpoint",
        "Azure portal and ADLS Gen2 ACLs",
        "Azure Databricks account or workspace console",
        "Azure Databricks SQL and Catalog Explorer",
        "Azure portal, ADLS Gen2, and Azure Databricks Catalog Explorer",
        "Azure Databricks SQL Warehouses and Compute",
        "Azure Databricks SQL or Catalog Explorer",
        "Azure Databricks SQL or Catalog Explorer",
        "Customer VDI, Azure CLI, and Databricks CLI",
        "Customer approval process",
    ]
    for row, platform in zip(setup_rows, platforms, strict=True):
        row.insert(2, platform)
    setup = add_table(
        workbook,
        "Customer Setup",
        [
            "ID",
            "Order",
            "Platform",
            "Customer administrator",
            "Admin setup action",
            "Detailed Azure or Azure Databricks setup steps",
            "Values customer must replace",
            "Actual command or action used",
            "Expected result and handoff",
            "Customer evidence or link",
            "Status",
            "Completed by",
            "Completed date",
            "Notes",
        ],
        setup_rows,
        [10, 9, 38, 34, 42, 74, 52, 62, 52, 34, 20, 24, 18, 44],
        "CustomerSetupTable",
    )
    add_list(setup, 11, "A", len(STATUSES), len(setup_rows))
    add_status_colors(setup, 11, len(setup_rows))

    security_rows = [
        [f"SEC-{index:03d}", "", "", "", "", "", "", "", "", "", "", "", "", "Not started", "", ""]
        for index in range(1, 31)
    ]
    security = add_table(
        workbook,
        "Security Mapping",
        [
            "Mapping ID",
            "Source principal",
            "Principal type",
            "Source object ID",
            "Source membership",
            "Source role, grant, or deny",
            "Source object or scope",
            "Required target behavior",
            "Existing target principal",
            "Proposed target principal",
            "Target role or grant",
            "Target object or scope",
            "Disposition",
            "Status",
            "Approval evidence",
            "Notes",
        ],
        security_rows,
        [12, 34, 22, 38, 38, 38, 38, 42, 34, 34, 38, 38, 22, 20, 28, 38],
        "SecurityMappingTable",
    )
    add_list(security, 13, "B", len(SECURITY_DISPOSITIONS), len(security_rows))
    add_list(security, 14, "A", len(STATUSES), len(security_rows))
    add_status_colors(security, 14, len(security_rows))

    lists = workbook.create_sheet("Lists")
    lists.sheet_state = "hidden"
    for column, (heading, values) in enumerate(
        [("Status", STATUSES), ("Security disposition", SECURITY_DISPOSITIONS)],
        start=1,
    ):
        lists.cell(1, column, heading)
        for row, value in enumerate(values, start=2):
            lists.cell(row, column, value)

    workbook.active = 1
    return workbook


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook()
    workbook.save(OUTPUT)
    loaded = load_workbook(OUTPUT, read_only=False, data_only=False)
    if loaded.sheetnames != ["Start Here", "Customer Setup", "Security Mapping", "Lists"]:
        raise RuntimeError(f"Unexpected workbook sheets: {loaded.sheetnames}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
