from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def test_skill_frontmatter_and_names_are_valid():
    skills = sorted((ROOT / "skills").glob("*/SKILL.md"))
    assert skills
    for path in skills:
        text = path.read_text(encoding="utf-8")
        assert text.startswith("---\n")
        _, frontmatter, body = text.split("---", 2)
        metadata = yaml.safe_load(frontmatter)
        assert metadata["name"] == path.parent.name
        assert metadata["description"]
        for section in ("## Must", "## Prefer", "## Avoid", "## Examples"):
            assert section in body, f"{path}: missing {section}"


def test_plugin_manifest_references_existing_skills():
    manifest_path = ROOT / "plugins/databricks-skills/.github/plugin/plugin.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["name"] == "databricks-skills"
    assert manifest["mcpServers"] == {}
    for reference in manifest["skills"]:
        name = reference.removeprefix("./skills/")
        assert (ROOT / "skills" / name / "SKILL.md").is_file()


def test_supported_migration_skills_are_present():
    expected = {
        "synapse-to-azure-databricks": {"references", "templates", "docs", "scripts"},
        "pipelines-to-azure-databricks-workflows": {"docs"},
    }
    for skill_name, expected_paths in expected.items():
        skill_root = ROOT / "skills" / skill_name
        assert (skill_root / "SKILL.md").is_file()
        for path_name in expected_paths:
            assert (skill_root / path_name).is_dir()


def test_eval_specs_have_required_shape():
    specs = sorted((ROOT / "tests/evals").glob("*/eval.yaml"))
    assert specs
    for path in specs:
        document = yaml.safe_load(path.read_text(encoding="utf-8"))
        assert document["name"]
        assert document["scoring"]["threshold"] == 1.0
        assert document["stimuli"]
        for stimulus in document["stimuli"]:
            assert stimulus.get("prompt") or stimulus.get("turns")
            assert stimulus["constraints"]["max_turns"] > 0
            grader_types = {grader["type"] for grader in stimulus["graders"]}
            assert {"skill-invocation", "completed", "token-budget", "error-count"} <= grader_types


def test_independence_audit_passes():
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts/check_independence.py")],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def test_synapse_serverless_sql_mode_is_complete():
    skill_root = ROOT / "skills/synapse-to-azure-databricks"
    skill_text = (skill_root / "SKILL.md").read_text(encoding="utf-8")
    reference = skill_root / "references/synapse-serverless-sql.md"
    reference_text = reference.read_text(encoding="utf-8")

    assert "synapse-serverless-sql.md" in skill_text
    for term in (
        "OPENROWSET",
        "CETAS",
        "Unity Catalog",
        "storage credential",
        "external location",
        "SQL Warehouse",
        "Migration manifest",
        "Cutover and rollback",
        "Lakebridge-first decision gate",
        "Mapping-gap gate",
        "DACPAC intake",
        "Post-staging transformation strategy",
        "Fact/dimension pipelines and persisted outputs",
        "Lakeflow Declarative Pipeline",
        "Persisted outputs",
    ):
        assert term in reference_text

    for term in (
        "Tool-routing contract",
        "Phase communication contract",
        "assessment support does not imply conversion support",
        "governed fallback",
        "Lakebridge Assessment Status",
        "Lakebridge Transpiler Status",
        "Migration Inventory Status",
    ):
        assert term in skill_text

    for template in (
        "mapping-gap-report.md",
        "migration-plan.md",
        "object-inventory.csv",
        "parity-test-manifest.yml",
        "validation-report.md",
    ):
        assert (skill_root / "templates" / template).is_file()

    parity_manifest = yaml.safe_load(
        (skill_root / "templates/parity-test-manifest.yml").read_text(encoding="utf-8")
    )
    validation_report = (skill_root / "templates/validation-report.md").read_text(encoding="utf-8")
    deployment_validation = (skill_root / "references/deployment-validation.md").read_text(encoding="utf-8")
    pipeline_test = next(test for test in parity_manifest["tests"] if test["mode"] == "pipeline_contract")
    assert pipeline_test["pipeline_checks"] == [
        "idempotent_replay",
        "watermark_advancement",
        "partial_failure_restart",
        "duplicate_prevention",
        "fact_grain",
        "dimension_referential_integrity",
        "late_arriving_data",
        "serving_view_result",
    ]
    for text in (validation_report, deployment_validation):
        assert "pipeline_contract" in text

    lifecycle_states = [
        "inventoried",
        "assessed",
        "fallback-required",
        "approved",
        "converted",
        "deployed",
        "validated",
        "failed",
        "retained",
        "blocked",
        "rolled-back",
    ]
    mapping_report = (skill_root / "templates/mapping-gap-report.md").read_text(encoding="utf-8")

    def table_values(start_heading: str, end_heading: str) -> list[str]:
        section = mapping_report.split(start_heading, 1)[1].split(end_heading, 1)[0]
        rows = []
        for line in section.splitlines():
            if not line.strip().startswith("|"):
                continue
            first_cell = line.strip().strip("|").split("|", 1)[0].strip()
            if first_cell and set(first_cell) - {"-", ":"}:
                rows.append(first_cell)
        return rows[1:]

    assert table_values("## Compatibility classification summary", "## Migration disposition summary") == [
        "direct",
        "rewrite",
        "redesign",
        "retain",
        "blocker",
    ]
    dispositions = ["migrate", "retain", "park", "retire"]
    assert table_values("## Migration disposition summary", "## Workflow status summary") == dispositions
    assert table_values("## Workflow status summary", "## Gaps") == lifecycle_states
    for state in lifecycle_states:
        assert f"`{state}`" in skill_text

    with (skill_root / "templates/object-inventory.csv").open(encoding="utf-8", newline="") as file:
        inventory_rows = list(csv.reader(file))
    disposition_index = inventory_rows[0].index("disposition")
    inventory_dispositions = inventory_rows[1][disposition_index].strip("<>").split("|")
    assert inventory_dispositions == dispositions
    migration_plan = (skill_root / "templates/migration-plan.md").read_text(encoding="utf-8")
    assert "migrate/retain/park/retire" in migration_plan
    assert "Use Lakeflow Declarative Pipelines" in reference_text
    assert "when it is declarative" in reference_text
    assert "Use Databricks Jobs for imperative orchestration" in reference_text
    conversion_text = (skill_root / "references/conversion.md").read_text(
        encoding="utf-8"
    )
    for term in (
        "Lakebridge conversion gate",
        "databricks labs lakebridge transpile --help",
        "databricks labs lakebridge install-transpile",
        "Assessment support does not imply conversion support",
        "representative pilot",
    ):
        assert term.lower() in conversion_text.lower()


def test_synapse_workload_guidance_is_separated():
    skill_root = ROOT / "skills/synapse-to-azure-databricks"
    skill_text = (skill_root / "SKILL.md").read_text(encoding="utf-8")
    references = {
        "serverless": skill_root / "references/synapse-serverless-sql.md",
        "dedicated": skill_root / "references/synapse-dedicated-sql-pool.md",
        "spark": skill_root / "references/synapse-spark.md",
    }
    for path in references.values():
        assert path.is_file()
        text = path.read_text(encoding="utf-8")
        for section in ("## Must", "## Prefer", "## Avoid", "## Examples"):
            assert section in text

    assert "name: synapse-to-azure-databricks" in skill_text
    assert "synapse-serverless-sql.md" in skill_text
    assert "synapse-dedicated-sql-pool.md" in skill_text
    assert "synapse-spark.md" in skill_text
    assert "staged-bronze-migration.md" in skill_text
    assert "Tool-routing contract" in skill_text
    assert "OPENROWSET" in references["serverless"].read_text(encoding="utf-8")
    assert "ROUND_ROBIN" in references["dedicated"].read_text(encoding="utf-8")
    assert "mssparkutils" in references["spark"].read_text(encoding="utf-8")

    docs = skill_root / "docs"
    for name in (
        "README.md",
        "skill-routing.md",
        "skill-routing.drawio",
        "prerequisites.md",
        "prerequisites.drawio",
        "migration-approach.md",
        "migration-approach.drawio",
        "synapse-to-azure-databricks-prerequisites.xlsx",
    ):
        path = docs / name
        assert path.is_file()
        if path.suffix in {".md", ".drawio"}:
            text = path.read_text(encoding="utf-8")
            assert "customer/" not in text.lower()
            assert "best buy" not in text.lower()
        if path.suffix == ".drawio":
            document = ET.parse(path)
            assert document.getroot().tag == "mxfile"

    workbook = load_workbook(
        docs / "synapse-to-azure-databricks-prerequisites.xlsx",
        read_only=True,
        data_only=False,
    )
    assert workbook.sheetnames == [
        "Start Here",
        "Customer Setup",
        "Security Mapping",
        "Lists",
    ]
    actions = workbook["Customer Setup"]
    headers = [cell.value for cell in next(actions.iter_rows(max_row=1))]
    assert "Platform" in headers
    assert "Detailed Azure or Azure Databricks setup steps" in headers
    assert "Actual command or action used" in headers
    assert "Expected result and handoff" in headers
    assert "Customer evidence or link" in headers
    assert actions.max_row - 1 >= 15
    action_text = " ".join(
        str(cell.value or "")
        for row in actions.iter_rows(min_row=2)
        for cell in row
    )
    for term in (
        "Microsoft Entra admin center",
        "Azure portal - Access control (IAM)",
        "Azure Databricks account or workspace console",
        "databricks labs install lakebridge",
        "databricks labs lakebridge --help",
        "Python 3.10.1-3.14.x",
        "Java 21+",
        "CREATE CATALOG",
        "CREATE EXTERNAL TABLE",
        "CAN USE",
        "databricks auth login",
        "Approve handoff",
    ):
        assert term.lower() in action_text.lower()
    security = workbook["Security Mapping"]
    for row in security.iter_rows(min_row=2, values_only=True):
        assert all(value is None for value in row[1:13])
    prerequisite_text = (docs / "prerequisites.md").read_text(encoding="utf-8")
    for term in (
        "customer-administrator setup",
        "separate named migration user",
        "GitHub Copilot",
        "customer VDI",
        "Install and verify Lakebridge",
        "databricks labs install lakebridge",
        "must stop and ask the user",
        "At every phase transition",
        "does not create workspaces",
        "After technical schema validation",
        "revoke temporary migration access",
    ):
        assert term.lower() in prerequisite_text.lower()

    smoke_eval = (
        ROOT / "tests/evals/synapse-to-azure-databricks/eval.yaml"
    ).read_text(encoding="utf-8")
    full_eval = (
        ROOT / "tests/evals/synapse-to-azure-databricks-fulleval/eval.yaml"
    ).read_text(encoding="utf-8")
    for text in (smoke_eval, full_eval):
        assert "Phase 1" in text
        assert "Lakebridge Assessment Status" in text
        assert "Lakebridge Transpiler Status" in text
        assert "Migration Inventory Status" in text
        assert "governed fallback" in text


def test_staged_bronze_scenario_uses_caf_controls():
    skill_root = ROOT / "skills/synapse-to-azure-databricks"
    skill_text = (skill_root / "SKILL.md").read_text(encoding="utf-8")
    references = {
        path.name: path.read_text(encoding="utf-8")
        for path in (skill_root / "references").glob("staged-bronze-*.md")
    }
    expected = {
        "staged-bronze-migration.md",
        "staged-bronze-external-registration.md",
        "staged-bronze-transformation-migration.md",
        "staged-bronze-managed-upgrade.md",
        "staged-bronze-adf-ingestion-migration.md",
        "staged-bronze-validation-cutover.md",
    }
    assert set(references) == expected
    for trigger in (
        "OSA migration",
        "external Bronze registration",
        "UC Upgrade",
        "ADF ingestion migration",
        "staged Bronze modernisation",
    ):
        assert trigger in skill_text
    profile = references["staged-bronze-migration.md"]
    for term in (
        "required Phase 1 Excel workbook and JSON manifest",
        "Lakebridge Assessment Status",
        "Lakebridge Transpiler Status",
        "Migration Inventory Status",
        "stage `N-1` is `PASS`",
        "writer-ownership transfer",
        "source retirement or file cleanup",
    ):
        assert term in profile
    assert "Reject direct `SET MANAGED` for Parquet" in references[
        "staged-bronze-managed-upgrade.md"
    ]
    assert "Keep schedules paused" in references[
        "staged-bronze-validation-cutover.md"
    ]

    smoke_eval = (
        ROOT / "tests/evals/synapse-to-azure-databricks/eval.yaml"
    ).read_text(encoding="utf-8")
    full_eval = (
        ROOT / "tests/evals/synapse-to-azure-databricks-fulleval/eval.yaml"
    ).read_text(encoding="utf-8")
    for text in (smoke_eval, full_eval):
        assert "staged Bronze" in text
        assert "writer ownership" in text


def test_plugin_build_copies_synapse_skill_assets():
    result = subprocess.run(
        [sys.executable, str(ROOT / "build/build_plugins.py"), "--check"],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    copied = ROOT / "plugins/databricks-skills/skills/synapse-to-azure-databricks/references/synapse-serverless-sql.md"
    assert copied.is_file()
    copied_templates = ROOT / "plugins/databricks-skills/skills/synapse-to-azure-databricks/templates"
    assert (copied_templates / "mapping-gap-report.md").is_file()
    assert (copied_templates / "parity-test-manifest.yml").is_file()
    copied_references = ROOT / "plugins/databricks-skills/skills/synapse-to-azure-databricks/references"
    assert (copied_references / "synapse-dedicated-sql-pool.md").is_file()
    assert (copied_references / "synapse-spark.md").is_file()
    assert (copied_references / "staged-bronze-migration.md").is_file()
    copied_docs = ROOT / "plugins/databricks-skills/skills/synapse-to-azure-databricks/docs"
    assert (copied_docs / "README.md").is_file()
    assert (copied_docs / "skill-routing.drawio").is_file()
    assert (copied_docs / "prerequisites.drawio").is_file()
    assert (copied_docs / "migration-approach.drawio").is_file()
    assert (copied_docs / "synapse-to-azure-databricks-prerequisites.xlsx").is_file()
    copied_scripts = ROOT / "plugins/databricks-skills/skills/synapse-to-azure-databricks/scripts"
    assert (copied_scripts / "create_migration_inventory.py").is_file()


def test_synapse_inventory_generator_creates_governed_artifacts():
    skill_root = ROOT / "skills/synapse-to-azure-databricks"
    generator = skill_root / "scripts/create_migration_inventory.py"
    assert generator.is_file()

    headers = next(csv.reader((skill_root / "templates/object-inventory.csv").open(encoding="utf-8")))
    row = {header: "" for header in headers}
    row.update(
        {
            "source_database": "sales",
            "source_schema": "dbo",
            "source_object": "orders",
            "source_type": "external_table",
            "definition_file": "exports/orders.sql",
            "definition_sha256": "a" * 64,
            "dependencies": "sales.dbo.customers",
            "source_identity": "synapse://workspace/sales/dbo/orders",
            "classification": "direct",
            "disposition": "migrate",
            "tool_route": "lakebridge",
            "target_catalog": "main",
            "target_schema": "sales",
            "target_object": "orders",
            "target_type": "external_table",
            "owner": "data-platform",
            "validation_status": "NOT RUN",
            "scenario_profile": "staged-bronze-modernisation",
            "scenario_stage": "1",
            "stage_gate_status": "not-started",
            "writer_owner": "adf-owned",
            "data_action": "register-existing",
            "original_location": "abfss://bronze@account.dfs.core.windows.net/orders",
            "managed_target": "not-applicable",
            "rollback_boundary": "metadata-only registration",
            "retention_decision": "pending",
            "cleanup_decision": "retain",
        }
    )

    with tempfile.TemporaryDirectory() as directory:
        temp = Path(directory)
        inventory_csv = temp / "object-inventory.csv"
        with inventory_csv.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=headers)
            writer.writeheader()
            writer.writerow(row)
        workbook_path = temp / "synapse-to-databricks-migration-inventory.xlsx"
        manifest_path = temp / "synapse-to-databricks-migration-inventory.manifest.json"
        result = subprocess.run(
            [
                sys.executable,
                str(generator),
                "--inventory-csv",
                str(inventory_csv),
                "--output",
                str(workbook_path),
                "--manifest",
                str(manifest_path),
            ],
            cwd=ROOT,
            text=True,
            capture_output=True,
            check=False,
        )
        assert result.returncode == 0, result.stdout + result.stderr
        workbook = load_workbook(workbook_path, data_only=False)
        assert workbook.sheetnames == [
            "Summary",
            "Migration Status",
            "Target Objects",
            "Dependencies",
            "Decisions and Approvals",
            "Evidence",
            "Metadata and Lists",
        ]
        assert workbook["Migration Status"].max_row == 2
        assert workbook["Migration Status"]["A2"].value == row["source_identity"]
        status_headers = [
            cell.value for cell in workbook["Migration Status"][1]
        ]
        status_row = {
            header: workbook["Migration Status"].cell(2, index + 1).value
            for index, header in enumerate(status_headers)
        }
        assert status_row["Lakebridge Assessment Status"] == "not-run"
        assert status_row["Lakebridge Transpiler Status"] == "not-evaluated"
        assert status_row["Migration Inventory Status"] == "inventoried"
        assert status_row["Scenario Profile"] == "staged-bronze-modernisation"
        assert status_row["Writer Ownership"] == "adf-owned"
        assert workbook["Target Objects"]["C2"].value == "orders"
        assert workbook["Dependencies"]["B2"].value == "sales.dbo.customers"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        assert len(manifest["items"]) == 1
        assert manifest["items"][0]["stable_source_id"] == row["source_identity"]
        assert manifest["items"][0]["lakebridge_assessment_status"] == "not-run"
        assert manifest["items"][0]["lakebridge_transpiler_status"] == "not-evaluated"
        assert manifest["items"][0]["migration_inventory_status"] == "inventoried"
        assert manifest["schema_version"] == 2
        assert manifest["items"][0]["scenario"]["stage"] == "1"
        assert manifest["items"][0]["scenario"]["cleanup_decision"] == "retain"

        invalid_cases = (
            ("scenario_stage", "", "scenario_stage"),
            ("stage_gate_status", "unknown", "stage_gate_status"),
            ("writer_owner", "unknown", "writer_owner"),
            ("data_action", "unknown", "data_action"),
            ("original_location", "", "original_location"),
            ("scenario_profile", "unknown-profile", "scenario_profile"),
        )
        for column, value, error_field in invalid_cases:
            invalid_row = row.copy()
            invalid_row[column] = value
            with inventory_csv.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=headers)
                writer.writeheader()
                writer.writerow(invalid_row)
            invalid_workbook = temp / f"invalid-{column}.xlsx"
            invalid_manifest = temp / f"invalid-{column}.json"
            result = subprocess.run(
                [
                    sys.executable,
                    str(generator),
                    "--inventory-csv",
                    str(inventory_csv),
                    "--output",
                    str(invalid_workbook),
                    "--manifest",
                    str(invalid_manifest),
                ],
                cwd=ROOT,
                text=True,
                capture_output=True,
                check=False,
            )
            assert result.returncode != 0
            assert f"row 2" in result.stderr
            assert error_field in result.stderr
            assert not invalid_workbook.exists()
            assert not invalid_manifest.exists()


def test_pipeline_skill_extends_serverless_inventory():
    skill_text = (
        ROOT / "skills/pipelines-to-azure-databricks-workflows/SKILL.md"
    ).read_text(encoding="utf-8")
    for term in (
        "Assumed platform baseline",
        "Migration Status",
        "Target Objects",
        "Pipeline Dependencies",
        "Pipeline Summary",
        "Lakebridge Assessment Status",
        "Lakebridge Transpiler Status",
        "Migration Inventory Status",
        "reconciled",
        "instead of creating a competing workbook",
    ):
        assert term in skill_text
    assert "### 0." not in skill_text
    assert "Create and maintain the `Pre Reqs`" not in skill_text
